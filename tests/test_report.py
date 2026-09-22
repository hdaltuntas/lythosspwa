"""
Tests for report.py: HTML content, the reportlab PDF, and DOCX if python-docx
is installed.
"""
import copy

import pytest

from lythosspwa import report
from lythosspwa.analysis_engine import AnalysisEngine, RetainingWall
from lythosspwa.beam_spring import BeamSpringAnalysis
from lythosspwa.config import DEFAULT_CONFIG


@pytest.fixture(scope="module")
def case():
    w = RetainingWall(copy.deepcopy(DEFAULT_CONFIG))
    e = AnalysisEngine(w); e.run()
    bs = BeamSpringAnalysis(w, e).run()
    return w, e, bs


@pytest.fixture(scope="module")
def cantilever():
    cfg = copy.deepcopy(DEFAULT_CONFIG)
    cfg['analysis_options'].update({'anchors': [], 'is_seismic': False})
    cfg['geometry']['excavation_depth_H'] = 4.0
    w = RetainingWall(cfg)
    e = AnalysisEngine(w); e.run()
    return w, e


def test_figures_rendered(case):
    w, e, bs = case
    figs = report.render_figures(w, e, bs, "en", dpi=60)
    assert set(figs) == set(report.FIGURE_KEYS)
    assert all(png.startswith(b"\x89PNG") for png in figs.values())
    figs = report.render_figures(w, e, None, "en", dpi=60)
    assert "beam_spring" not in figs


@pytest.mark.parametrize("lang", ["en", "tr"])
def test_html_contains_key_numbers(case, lang):
    w, e, bs = case
    figs = report.render_figures(w, e, bs, lang, dpi=40)
    html_text = report.build_html(w, e, bs, lang, figs)
    assert f"{e.d_required:,.2f}" in html_text and f"{e.d_design:,.2f}" in html_text
    assert f"{bs['m_max_abs']:,.1f}" in html_text
    for a in w.anchors:
        assert f"{bs['anchor_report'][a['depth']]['T_axial']:,.1f}" in html_text
    assert html_text.count("data:image/png;base64,") == len(figs)
    T = report.TEXTS[lang]
    for key in ("sec_inputs", "sec_le", "sec_bs", "sec_figs", "sec_notes"):
        assert T[key] in html_text


def test_html_without_beam_spring_and_cantilever(cantilever):
    w, e = cantilever
    figs = report.render_figures(w, e, None, "tr", dpi=40)
    html_text = report.build_html(w, e, None, "tr", figs)
    assert report.TEXTS["tr"]["sec_bs"] not in html_text
    assert report.TEXTS["tr"]["toe_R"] in html_text


@pytest.mark.parametrize("lang", ["en", "tr"])
def test_pdf_export(case, tmp_path, lang):
    w, e, bs = case
    path = tmp_path / "r.pdf"
    report.export_report(str(path), w, e, bs, lang)
    data = path.read_bytes()
    assert data.startswith(b"%PDF") and len(data) > 100_000


def test_pdf_font_covers_the_alphabets_the_report_uses():
    """Turkish and the Greek symbols need a Unicode font; Helvetica has neither."""
    from lythosspwa import pdf
    regular, bold = pdf.fonts()
    assert regular != "Helvetica", "no Unicode font was registered for the PDF"
    from reportlab.pdfbase import pdfmetrics
    for name in (regular, bold):
        face = pdfmetrics.getFont(name).face
        # Turkish, the Greek letters of soil mechanics, and the symbols the
        # tables are written with.
        for character in "ışğİÖÜçÇ" + "φγδβψ" + "°³ₛ√≥×′":
            assert face.charToGlyph.get(ord(character)), \
                f"{name} has no glyph for {character}"


def test_pdf_walker_renders_headings_tables_and_figures():
    """The walker must turn the report's markup into flowables, not drop it."""

    from lythosspwa import pdf
    html_text = ("<h1>Title</h1><p class='meta'>meta</p>"
                 "<h2 style='page-break-before:always'>Section</h2>"
                 "<h3>Sub</h3>"
                 "<table width='100%'><tr><th width='60%'>a</th><th width='40%'>b</th></tr>"
                 "<tr><td>1</td><td><span class=\"ok\">OK</span></td></tr></table>"
                 "<ul><li>note</li></ul>"
                 "<p><img src='fig://moment' width='640'></p>")
    figures = {"moment": _one_pixel_png()}
    story = pdf.flowables(html_text, figures, pdf._styles(), 400, 700)
    kinds = [type(item).__name__ for item in story]
    assert kinds.count("Paragraph") >= 4
    assert "Table" in kinds and "PageBreak" in kinds
    assert any("Image" in kind or "KeepTogether" in kind for kind in kinds)


def _one_pixel_png() -> bytes:
    """The smallest valid PNG, so the walker has a figure to place."""
    import base64
    return base64.b64decode(
        b"iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mP8z8BQDwAEhQGAhKmMIQ"
        b"AAAABJRU5ErkJggg==")


def test_html_export(case, tmp_path):
    w, e, bs = case
    path = tmp_path / "r.html"
    report.export_report(str(path), w, e, bs, "tr")
    assert "<html" in path.read_text(encoding="utf-8")[:200]


def test_docx_export(case, tmp_path):
    pytest.importorskip("docx")
    w, e, bs = case
    path = tmp_path / "r.docx"
    report.export_report(str(path), w, e, bs, "en")
    import docx
    d = docx.Document(str(path))
    headings = [p.text for p in d.paragraphs if p.style.name.startswith(("Heading", "Title"))]
    assert report.TEXTS["en"]["sec_le"] in headings
    assert len(d.tables) >= 10 and len(d.inline_shapes) == len(report.FIGURE_KEYS)
