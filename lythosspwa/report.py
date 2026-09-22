"""
Calculation report for Lythos SPWA.

The report is assembled as HTML (inputs, LE results, beam-spring results,
checks, figures, warnings) and exported as
  * PDF   — through reportlab (see `lythosspwa.pdf`)
  * HTML  — single self-contained file (figures embedded as base64)
  * DOCX  — optional, needs `python-docx`

There is one assembly, `build_html()`, so all three formats say the same thing.

Figures are rendered with plotting.Plotter on off-screen Matplotlib figures.
"""

import base64
import datetime
import html
import io
from typing import Any, Dict, List, Optional

from matplotlib.figure import Figure
from scipy.stats import norm

from . import study_plots
from .config import APP_NAME, APP_VERSION, TRANSLATIONS
from .plotting import Plotter
from .study import OUTPUTS

FIGURE_KEYS = ["net_pressure", "earth_pressure", "water_pressure",
               "shear", "moment", "deflection", "beam_spring"]

TEXTS = {
    "en": {
        "title": f"{APP_NAME} — Calculation Report",
        "date": "Date", "analyst": "Analyst", "software": "Software",
        "sec_inputs": "1. Input data", "sec_geom": "1.1 Geometry, loads and water",
        "sec_soil": "1.2 Soil profile", "sec_anchors": "1.3 Anchors",
        "sec_struct": "1.4 Wall section", "sec_opts": "1.5 Analysis options and factors",
        "sec_le": "2. Limit-equilibrium analysis (free-earth support)",
        "sec_le_embed": "2.1 Embedment", "sec_le_anchor": "2.2 Anchor forces",
        "sec_le_summary": "2.3 Internal forces and displacements (over H + D_req)",
        "sec_checks": "2.4 Design checks",
        "sec_bs": "3. Beam-spring (Winkler) analysis", "sec_bs_ks": "3.1 Subgrade reaction moduli",
        "sec_bs_stages": "3.2 Construction stages", "sec_bs_anchor": "3.3 Anchor forces (final stage)",
        "sec_bs_checks": "3.4 Design checks (final stage)",
        "sec_figs": "4. Figures", "sec_warn": "5. Warnings", "sec_notes": "6. Method notes",
        "parameter": "Parameter", "value": "Value", "unit": "Unit",
        "H": "Excavation depth H", "beta": "Backfill slope β", "alpha": "Dredge line slope α",
        "delta": "Wall friction angle δ", "q": "Surcharge", "hw_a": "Water level, retained side",
        "hw_p": "Water level, excavation side", "gw": "Unit weight of water",
        "layer": "Layer", "thickness": "Thickness", "gamma": "γ", "gamma_sat": "γsat", "phi": "φ",
        "c": "c", "ks": "kₛ", "ks_method": "kₛ method",
        "anchor": "Anchor", "depth": "Depth", "angle": "Angle", "EA": "EA", "free_length": "Free length",
        "spacing": "Spacing", "prestress": "Lock-off load", "k_h": "k_h",
        "manufacturer": "Manufacturer", "section": "Section", "I": "Moment of inertia I",
        "W": "Section modulus W", "E": "Young's modulus E", "grade": "Steel grade", "fy": "Yield strength f_y",
        "f_allow": "Allowable bending stress f_y / FS", "EI": "Flexural stiffness EI",
        "seismic": "Seismic analysis", "kh": "kh", "kv": "kv", "sub_theta": "Submerged θ correction",
        "hydro": "Westergaard hydrodynamic pressure", "fs_phi": "Partial factor on tan φ",
        "fs_c": "Partial factor on c", "fs_b": "Factor on bending (f_y / FS)",
        "emb_factor": "Embedment increase factor", "defl_code": "Deflection criterion",
        "bs_enabled": "Beam-spring analysis", "bs_staged": "Staged construction", "bs_overdig": "Overdig",
        "bs_water": "Water in excavation", "yes": "yes", "no": "no",
        "water_final": "as given (underwater dredging)", "water_dewatered": "dewatered until final stage",
        "d_req": "Theoretical embedment D_req", "d_design": "Design embedment D_design",
        "L": "Total wall length", "toe_R": "Toe reaction R (simplified method)",
        "closure": "Moment at theoretical toe (closure check)",
        "T_h": "T_h (horizontal)", "T_ax": "Axial per anchor", "V": "Vertical component",
        "quantity": "Quantity", "min": "Min", "max": "Max",
        "p": "Net pressure", "Vf": "Shear force", "M": "Bending moment", "rot": "Rotation", "defl": "Deflection",
        "check": "Check", "actual": "Actual", "allowable": "Allowable", "status": "Status",
        "stress": "Bending stress", "defl_check": "Deflection", "vertical": "Vertical equilibrium (indicative)",
        "ok": "OK", "notok": "NOT OK", "na": "n/a",
        "stage": "Stage", "description": "Description", "w_max": "w_max", "forces": "Anchor forces (kN/m)",
        "excavate": "Excavate to {val:.2f} m", "install": "Install anchor at {val:.2f} m",
        "embed_used": "Embedment used", "iterations": "Newton iterations",
        "mob_p": "Mobilised passive resistance", "mob_a": "Retained face at active limit",
        "compare": "Comparison LE vs beam-spring",
        "sec_study": "7. Parametric / reliability study", "sec_study_vars": "7.1 Variables",
        "sec_study_stats": "7.2 Statistics of the results", "sec_study_rel": "7.3 Reliability",
        "sec_study_sens": "7.4 Sensitivities (Spearman ρ)", "sec_study_figs": "7.5 Figures",
        "study_method": "Sampling method", "study_n": "Number of samples", "study_ok": "Successful analyses",
        "study_bs": "Beam-spring included", "study_unf": "Unfactored strengths",
        "mode": "Mode", "range": "Range", "dist": "Distribution", "mean": "Mean", "cov": "CoV",
        "output": "Output", "std": "Std", "limit_state": "Limit state", "n_fail": "Failures",
        "pf": "P_f", "pf_ci": "95 % CI", "beta_idx": "β", "beta_fosm": "β (mean/std of margin)",
        "variable": "Variable",
        "multi_note": "The free-earth method is statically determinate for one anchor only; for several anchors "
                      "the LE distribution is approximate and the beam-spring results should govern.",
        "notes": [
            "Earth pressures: Coulomb (static) / Mononobe-Okabe (seismic) for a vertical wall; cohesion term "
            "2c√K; tension cut-off on the active side.",
            "LE embedment from moment equilibrium about the toe (cantilever, simplified method) or the lowest "
            "anchor (free-earth). Internal forces at D_req; D_design = round-up(1.2·D_req).",
            "Below the water table the seismic inertia angle uses γsat/γ′ (restrained pore water). Westergaard "
            "hydrodynamic pressure 7/8·kh·γw·√(H_w·y) is applied to the free water in front of the wall.",
            "Beam-spring: Euler-Bernoulli beam on elastoplastic Winkler springs bounded by the active and passive "
            "limits, at-rest start (K₀ = 1 − sin φ); anchors are tension-only springs with k_h = EA/(L_free·s)·cos²α; "
            "springs keep their state between stages.",
            "Vertical check compares ΣT_h·tanα with the skin friction ∫(p_a + p_p)·tanδ on the embedded length; "
            "no end bearing — indicative only.",
        ],
    },
    "tr": {
        "title": f"{APP_NAME} — Hesap Raporu",
        "date": "Tarih", "analyst": "Hazırlayan", "software": "Yazılım",
        "sec_inputs": "1. Girdi verileri", "sec_geom": "1.1 Geometri, yükler ve su",
        "sec_soil": "1.2 Zemin profili", "sec_anchors": "1.3 Ankrajlar",
        "sec_struct": "1.4 Duvar kesiti", "sec_opts": "1.5 Analiz seçenekleri ve katsayılar",
        "sec_le": "2. Limit denge analizi (serbest zemin desteği)",
        "sec_le_embed": "2.1 Gömülme", "sec_le_anchor": "2.2 Ankraj kuvvetleri",
        "sec_le_summary": "2.3 İç kuvvetler ve deplasmanlar (H + D_req boyunca)",
        "sec_checks": "2.4 Tasarım kontrolleri",
        "sec_bs": "3. Kiriş-yay (Winkler) analizi", "sec_bs_ks": "3.1 Yatak katsayıları",
        "sec_bs_stages": "3.2 İmalat aşamaları", "sec_bs_anchor": "3.3 Ankraj kuvvetleri (son aşama)",
        "sec_bs_checks": "3.4 Tasarım kontrolleri (son aşama)",
        "sec_figs": "4. Şekiller", "sec_warn": "5. Uyarılar", "sec_notes": "6. Yöntem notları",
        "parameter": "Parametre", "value": "Değer", "unit": "Birim",
        "H": "Kazı derinliği H", "beta": "Dolgu şev açısı β", "alpha": "Tarama hattı şev açısı α",
        "delta": "Duvar sürtünme açısı δ", "q": "Sürşarj", "hw_a": "Su seviyesi, arka taraf",
        "hw_p": "Su seviyesi, kazı tarafı", "gw": "Suyun birim hacim ağırlığı",
        "layer": "Tabaka", "thickness": "Kalınlık", "gamma": "γ", "gamma_sat": "γdoy", "phi": "φ",
        "c": "c", "ks": "kₛ", "ks_method": "kₛ yöntemi",
        "anchor": "Ankraj", "depth": "Derinlik", "angle": "Açı", "EA": "EA", "free_length": "Serbest boy",
        "spacing": "Aralık", "prestress": "Kilitleme yükü", "k_h": "k_h",
        "manufacturer": "Üretici", "section": "Kesit", "I": "Atalet momenti I",
        "W": "Mukavemet momenti W", "E": "Elastisite modülü E", "grade": "Çelik sınıfı", "fy": "Akma dayanımı f_y",
        "f_allow": "İzin verilen eğilme gerilmesi f_y / FS", "EI": "Eğilme rijitliği EI",
        "seismic": "Sismik analiz", "kh": "kh", "kv": "kv", "sub_theta": "Batık θ düzeltmesi",
        "hydro": "Westergaard hidrodinamik basıncı", "fs_phi": "tan φ kısmi katsayısı",
        "fs_c": "c kısmi katsayısı", "fs_b": "Eğilme katsayısı (f_y / FS)",
        "emb_factor": "Gömülme artırma katsayısı", "defl_code": "Deplasman ölçütü",
        "bs_enabled": "Kiriş-yay analizi", "bs_staged": "Aşamalı imalat", "bs_overdig": "Kazı payı",
        "bs_water": "Kazıdaki su", "yes": "evet", "no": "hayır",
        "water_final": "girildiği gibi (su altı tarama)", "water_dewatered": "son aşamaya kadar susuzlaştırılmış",
        "d_req": "Teorik gömülme D_req", "d_design": "Tasarım gömülmesi D_design",
        "L": "Toplam duvar boyu", "toe_R": "Uç reaksiyonu R (basitleştirilmiş yöntem)",
        "closure": "Teorik uçta moment (kapanma kontrolü)",
        "T_h": "T_h (yatay)", "T_ax": "Ankraj başına eksenel", "V": "Düşey bileşen",
        "quantity": "Büyüklük", "min": "Min", "max": "Maks",
        "p": "Net basınç", "Vf": "Kesme kuvveti", "M": "Eğilme momenti", "rot": "Dönme", "defl": "Deplasman",
        "check": "Kontrol", "actual": "Oluşan", "allowable": "İzin verilen", "status": "Durum",
        "stress": "Eğilme gerilmesi", "defl_check": "Deplasman", "vertical": "Düşey denge (gösterge)",
        "ok": "UYGUN", "notok": "UYGUN DEĞİL", "na": "—",
        "stage": "Aşama", "description": "Açıklama", "w_max": "w_maks", "forces": "Ankraj kuvvetleri (kN/m)",
        "excavate": "{val:.2f} m'ye kazı", "install": "{val:.2f} m'de ankraj montajı",
        "embed_used": "Kullanılan gömülme", "iterations": "Newton iterasyonu",
        "mob_p": "Mobilize pasif direnç", "mob_a": "Arka yüzde aktif sınıra ulaşan kısım",
        "compare": "LE – kiriş-yay karşılaştırması",
        "sec_study": "7. Parametrik / güvenilirlik çalışması", "sec_study_vars": "7.1 Değişkenler",
        "sec_study_stats": "7.2 Sonuç istatistikleri", "sec_study_rel": "7.3 Güvenilirlik",
        "sec_study_sens": "7.4 Duyarlılıklar (Spearman ρ)", "sec_study_figs": "7.5 Şekiller",
        "study_method": "Örnekleme yöntemi", "study_n": "Örnek sayısı", "study_ok": "Başarılı analiz",
        "study_bs": "Kiriş-yay dahil", "study_unf": "Faktörsüz dayanımlar",
        "mode": "Mod", "range": "Aralık", "dist": "Dağılım", "mean": "Ortalama", "cov": "CoV",
        "output": "Çıktı", "std": "Std", "limit_state": "Sınır durumu", "n_fail": "Göçme",
        "pf": "P_f", "pf_ci": "%95 GA", "beta_idx": "β", "beta_fosm": "β (marj ort./std)",
        "variable": "Değişken",
        "multi_note": "Serbest zemin desteği yöntemi yalnız tek ankraj için statikçe belirlidir; çok ankrajlı "
                      "durumda LE dağılımı yaklaşıktır, kiriş-yay sonuçları esas alınmalıdır.",
        "notes": [
            "Zemin basınçları: düşey duvar için Coulomb (statik) / Mononobe-Okabe (sismik); kohezyon terimi 2c√K; "
            "aktif tarafta çekme kesmesi.",
            "LE gömülmesi uç (konsol, basitleştirilmiş yöntem) veya en alt ankraj (serbest zemin desteği) etrafında "
            "moment dengesinden. İç kuvvetler D_req'de; D_design = yukarı yuvarla(1.2·D_req).",
            "Su tablası altında sismik atalet açısı γdoy/γ′ ile hesaplanır (tutulan boşluk suyu). Duvar önündeki "
            "serbest suya Westergaard hidrodinamik basıncı 7/8·kh·γw·√(H_w·y) uygulanır.",
            "Kiriş-yay: aktif ve pasif sınırlarla sınırlı elastoplastik Winkler yayları üzerinde Euler-Bernoulli "
            "kirişi, sükûnet başlangıcı (K₀ = 1 − sin φ); ankrajlar yalnız çekme alan yaylar, k_h = EA/(L_serbest·s)·cos²α; "
            "yaylar aşamalar arasında durumlarını korur.",
            "Düşey kontrol ΣT_h·tanα'yı gömülü boydaki çevre sürtünmesi ∫(p_a + p_p)·tanδ ile karşılaştırır; "
            "uç direnci yok — yalnızca gösterge niteliğindedir.",
        ],
    },
}

_CSS = """
body { font-family: 'Segoe UI', Arial, sans-serif; font-size: 10pt; color: #222; }
h1 { font-size: 18pt; color: #1F4E79; margin-bottom: 2px; }
h2 { font-size: 13pt; color: #1F4E79; border-bottom: 1px solid #1F4E79; margin-top: 18px; }
h3 { font-size: 11pt; color: #333; margin-top: 12px; }
table { border-collapse: collapse; margin: 4px 0 8px 0; }
th { background: #E8EEF6; text-align: left; padding: 3px 6px; border: 1px solid #B8C4D6; font-size: 9pt; }
td { padding: 3px 6px; border: 1px solid #B8C4D6; font-size: 9pt; }
.ok { color: #1E8449; font-weight: bold; } .bad { color: #C0392B; font-weight: bold; }
.meta { color: #555; } .note { color: #555; font-size: 9pt; }
"""


def _esc(x) -> str:
    return html.escape(str(x))


def _f(x, nd=2) -> str:
    return f"{x:,.{nd}f}"


def _status(T, s: str) -> str:
    if s == "N/A":
        return T["na"]
    return f'<span class="ok">{T["ok"]}</span>' if s == "OK" else f'<span class="bad">{T["notok"]}</span>'


def _table(headers: List[str], rows: List[List[Any]], widths: Optional[List[int]] = None) -> str:
    out = ["<table width='100%'>"]
    if widths:
        cells = "".join(f"<th width='{w}%'>{_esc(h)}</th>" for h, w in zip(headers, widths))
    else:
        cells = "".join(f"<th>{_esc(h)}</th>" for h in headers)
    out.append("<tr>" + cells + "</tr>")
    for r in rows:
        out.append("<tr>" + "".join(f"<td>{c}</td>" for c in r) + "</tr>")
    out.append("</table>")
    return "\n".join(out)


def _kv_table(T, rows: List[List[Any]]) -> str:
    return _table([T["parameter"], T["value"], T["unit"]], rows, [50, 36, 14])


# ----------------------------------------------------------------------
def render_figures(wall, engine, bs, lang: str, keys=FIGURE_KEYS,
                   dpi: int = 130) -> Dict[str, bytes]:
    """Renders the report figures to PNG bytes (off-screen)."""
    plotter = Plotter(wall, engine, TRANSLATIONS[lang], bs)
    out = {}
    for key in keys:
        if key == 'beam_spring' and bs is None:
            continue
        fig = Figure(figsize=(10, 7), dpi=dpi)
        plotter.setup_plot(key, fig)
        buf = io.BytesIO()
        fig.savefig(buf, format='png', dpi=dpi, bbox_inches='tight')
        out[key] = buf.getvalue()
    return out


def render_study_figures(study, lang: str, dpi: int = 110) -> Dict[str, bytes]:
    """Study figures keyed 'study_<view>' (only the views that apply)."""
    if study is None or not study.rows:
        return {}
    L = TRANSLATIONS[lang]
    outs = study_plots.default_outputs(study)
    views = []
    if study.method == "oat":
        views.append(("study_oat", lambda f: study_plots.plot_oat(f, study, L)))
    else:
        views.append(("study_hist", lambda f: study_plots.plot_hist(f, study, L)))
        views.append(("study_scatter", lambda f: study_plots.plot_scatter(f, study, L, outs[0])))
        views.append(("study_tornado", lambda f: study_plots.plot_tornado(f, study, L, outs[0])))
    out = {}
    for key, fn in views:
        fig = Figure(figsize=(10, 7), dpi=dpi)
        fn(fig)
        buf = io.BytesIO()
        fig.savefig(buf, format='png', dpi=dpi, bbox_inches='tight')
        out[key] = buf.getvalue()
    return out


def _study_section(T, L, study, figures, img_src) -> List[str]:
    parts = [f"<h2 style='page-break-before:always'>{T['sec_study']}</h2>"]
    s = study.summary
    yn = lambda b: T["yes"] if b else T["no"]
    parts.append(_kv_table(T, [
        [T["study_method"], L[f"method_{study.method}"], ""],
        [T["study_n"], str(s.get('n_total', 0)), ""],
        [T["study_ok"], f"{s.get('n_ok', 0)} (LE) / {s.get('n_ok_bs', 0)} (Winkler)", ""],
        [T["study_bs"], yn(study.run_bs), ""],
        [T["study_unf"], yn(study.unfactored), ""],
    ]))
    parts.append(f"<h3>{T['sec_study_vars']}</h3>")
    rows = []
    for v in study.variables:
        if v.mode == "range":
            rows.append([_esc(v.label), T["range"], _f(v.vmin, 3), _f(v.vmax, 3), "—", "—", "—", v.n_points])
        else:
            rows.append([_esc(v.label), T["dist"], "—", "—", L[f"dist_{v.dist}"], _f(v.mean, 3), _f(v.cov, 3), "—"])
    parts.append(_table([T["variable"], T["mode"], T["min"], T["max"], T["dist"], T["mean"], T["cov"],
                         "n"], rows, [28, 12, 10, 10, 14, 10, 8, 8]))
    parts.append(f"<h3>{T['sec_study_stats']}</h3>")
    rows = [[L.get(f"out_{k}", k), st['n'], _f(st['mean']), _f(st['std']), _f(st['p5']), _f(st['p50']),
             _f(st['p95']), u] for (k, _, u) in OUTPUTS if k in s.get('stats', {})
            for st in [s['stats'][k]]]
    parts.append(_table([T["output"], "n", T["mean"], T["std"], "P5", "P50", "P95", T["unit"]], rows,
                        [22, 8, 12, 12, 12, 12, 12, 10]))
    if s.get('reliability'):
        parts.append(f"<h3>{T['sec_study_rel']}</h3>")
        rows = []
        for name, r in s['reliability'].items():
            if r['beta'] in (float('inf'), float('-inf')):
                beta = f"> {_f(float(-norm.ppf(3 / max(r['n'], 1))))}"
            else:
                beta = _f(r['beta'])
            rows.append([name, r['n'], r['n_fail'], f"{r['pf']:.3g}", f"[{r['pf_lo']:.2g}, {r['pf_hi']:.2g}]",
                         beta, _f(r['beta_fosm'])])
        parts.append(_table([T["limit_state"], "n", T["n_fail"], T["pf"], T["pf_ci"], T["beta_idx"], T["beta_fosm"]],
                            rows, [18, 8, 10, 12, 20, 12, 20]))
    if s.get('spearman'):
        parts.append(f"<h3>{T['sec_study_sens']}</h3>")
        outs = [k for k in study_plots.default_outputs(study) if k in s['spearman']]
        rows = [[_esc(v.label)] + [f"{s['spearman'][k].get(v.path, float('nan')):+.2f}" for k in outs]
                for v in study.variables]
        parts.append(_table([T["variable"]] + [L.get(f"out_{k}", k) for k in outs], rows))
    parts.append(f"<h3>{T['sec_study_figs']}</h3>")
    for key in figures:
        if key.startswith("study_"):
            parts.append(f"<p><img src='{img_src(key)}' width='640'></p>")
    return parts


def build_html(wall, engine, bs: Optional[Dict[str, Any]], lang: str,
               figures: Dict[str, bytes], img_src=None, study=None) -> str:
    """
    Builds the report HTML. img_src(key) -> value for the <img src> attribute;
    default embeds base64 data URIs (browser, and the self-contained HTML
    file). The PDF writer passes a `fig://<key>` mapper instead and resolves
    the keys against the figure dictionary itself.
    """
    T = TEXTS[lang]
    L = TRANSLATIONS[lang]
    cfg = wall.config
    res = engine.results
    if img_src is None:
        def img_src(key):
            return "data:image/png;base64," + base64.b64encode(figures[key]).decode()

    parts = [f"<html><head><meta charset='utf-8'><style>{_CSS}</style></head><body>"]
    parts.append(f"<h1>{_esc(cfg['project_info'].get('title', ''))}</h1>")
    parts.append(f"<p class='meta'>{_esc(T['title'])}<br>"
                 f"{T['date']}: {datetime.date.today().isoformat()} &nbsp;|&nbsp; "
                 f"{T['analyst']}: {_esc(cfg['project_info'].get('analyst', ''))} &nbsp;|&nbsp; "
                 f"{T['software']}: {APP_NAME} v{APP_VERSION}</p>")

    # ---------------- 1. inputs
    parts.append(f"<h2>{T['sec_inputs']}</h2>")
    parts.append(f"<h3>{T['sec_geom']}</h3>")
    geo, loads = cfg['geometry'], cfg['loads']
    parts.append(_kv_table(T, [
        [T["H"], _f(geo['excavation_depth_H']), "m"],
        [T["beta"], _f(geo['backfill_slope_beta'], 1), "°"],
        [T["alpha"], _f(geo['dredge_line_slope_alpha'], 1), "°"],
        [T["delta"], _f(geo['wall_friction_delta'], 1), "°"],
        [T["q"], _f(loads['surcharge_load'], 1), "kPa"],
        [T["hw_a"], _f(loads['water_level_active']), "m"],
        [T["hw_p"], _f(loads['water_level_passive']), "m"],
        [T["gw"], _f(cfg['constants']['gamma_water']), "kN/m³"],
    ]))

    parts.append(f"<h3>{T['sec_soil']}</h3>")
    rows = []
    for i, lay in enumerate(cfg['soil_profile'], 1):
        method = lay.get('k_s_method', 'manual')
        ks_txt = (_f(lay.get('k_s', 0), 0) if method == 'manual'
                  else f"E_M={_f(lay.get('E_M', 0), 1)} MPa, α={lay.get('alpha', 0.33):.2f}")
        rows.append([i, _esc(lay['name']), _f(lay['thickness']), _f(lay['gamma'], 1),
                     _f(lay['gamma_sat'], 1), _f(lay['phi'], 1), _f(lay['cohesion'], 1),
                     L["ks_" + method], ks_txt])
    parts.append(_table(["#", T["layer"], f"{T['thickness']} (m)", f"{T['gamma']} (kN/m³)",
                         f"{T['gamma_sat']} (kN/m³)", f"{T['phi']} (°)", f"{T['c']} (kPa)",
                         T["ks_method"], f"{T['ks']} (kN/m³)"], rows,
                        [4, 19, 10, 10, 10, 9, 9, 13, 16]))

    parts.append(f"<h3>{T['sec_anchors']}</h3>")
    if wall.anchors:
        rows = [[i, _f(a['depth']), _f(a['angle'], 1), _f(a['EA'], 0), _f(a['free_length']),
                 _f(a['spacing']), _f(a['prestress'], 1), _f(a['stiffness'], 0)]
                for i, a in enumerate(wall.anchors, 1)]
        parts.append(_table(["#", f"{T['depth']} (m)", f"{T['angle']} (°)", f"{T['EA']} (kN)",
                             f"{T['free_length']} (m)", f"{T['spacing']} (m)",
                             f"{T['prestress']} (kN)", f"{T['k_h']} (kN/m/m)"], rows,
                        [4, 11, 10, 15, 13, 11, 17, 19]))
    else:
        parts.append(f"<p>{_esc(L['no_anchors'].strip())}</p>")

    parts.append(f"<h3>{T['sec_struct']}</h3>")
    st = cfg['structural_properties']
    parts.append(_kv_table(T, [
        [T["manufacturer"], _esc(st['selected_manufacturer']), ""],
        [T["section"], _esc(wall.selected_section_model), ""],
        [T["I"], _f(wall.i_moment * 1e8, 0), "cm⁴/m"],
        [T["W"], _f(wall.w_modulus * 1e6, 0), "cm³/m"],
        [T["E"], _f(wall.e_modulus / 1000, 0), "MPa"],
        [T["EI"], _f(wall.ei, 0), "kNm²/m"],
        [T["grade"], _esc(wall.selected_steel_grade), ""],
        [T["fy"], _f(wall.fy_yield / 1000, 0), "MPa"],
        [T["f_allow"], _f(wall.f_allowable / 1000, 1), "MPa"],
    ]))

    parts.append(f"<h3>{T['sec_opts']}</h3>")
    opts, fac = cfg['analysis_options'], cfg['factors']
    yn = lambda b: T["yes"] if b else T["no"]
    bso = opts.get('beam_spring', {})
    parts.append(_kv_table(T, [
        [T["seismic"], yn(wall.is_seismic), ""],
        [T["kh"], _f(wall.kh, 3), ""], [T["kv"], _f(wall.kv, 3), ""],
        [T["sub_theta"], yn(wall.is_seismic and wall.submerged_theta), ""],
        [T["hydro"], yn(wall.is_seismic and wall.hydrodynamic), ""],
        [T["fs_phi"], _f(fac['FS_friction_angle']), ""],
        [T["fs_c"], _f(fac['FS_cohesion']), ""],
        [T["fs_b"], _f(fac['FS_bending']), ""],
        [T["emb_factor"], _f(fac['embedment_increase_factor']), ""],
        [T["defl_code"], _esc(opts['deflection_check_code']), ""],
        [T["bs_enabled"], yn(bs is not None), ""],
        [T["bs_staged"], yn(bso.get('staged', True)), ""],
        [T["bs_overdig"], _f(bso.get('overdig', 0.5)), "m"],
        [T["bs_water"], T["water_" + bso.get('water_mode', 'final')], ""],
    ]))

    # ---------------- 2. LE results
    parts.append(f"<h2 style='page-break-before:always'>{T['sec_le']}</h2>")
    parts.append(f"<h3>{T['sec_le_embed']}</h3>")
    rows = [[T["d_req"], _f(engine.d_required), "m"],
            [T["d_design"], _f(engine.d_design), "m"],
            [T["L"], _f(wall.h + engine.d_design), "m"],
            [T["closure"], _f(res['toe_moment']), "kNm/m"]]
    if engine.is_cantilever:
        rows.append([T["toe_R"], _f(abs(res['toe_shear']), 1), "kN/m"])
    parts.append(_kv_table(T, rows))

    parts.append(f"<h3>{T['sec_le_anchor']}</h3>")
    if wall.anchors:
        rows = [[i, _f(a['depth']), _f(rep['T_h'], 1), _f(rep['T_axial'], 1), _f(rep['V'], 1)]
                for i, (a, rep) in enumerate(
                    ((a, res['anchor_report'][a['depth']]) for a in wall.anchors), 1)]
        parts.append(_table(["#", f"{T['depth']} (m)", f"{T['T_h']} (kN/m)",
                             f"{T['T_ax']} (kN)", f"{T['V']} (kN/m)"], rows))
        if len(wall.anchors) > 1:
            parts.append(f"<p class='note'>{T['multi_note']}</p>")
    else:
        parts.append(f"<p>{_esc(L['no_anchors'].strip())}</p>")

    parts.append(f"<h3>{T['sec_le_summary']}</h3>")
    parts.append(_table([T["quantity"], T["min"], T["max"], T["unit"]], [
        [T["p"], _f(res['p_min']), _f(res['p_max']), "kPa"],
        [T["Vf"], _f(res['v_min']), _f(res['v_max']), "kN/m"],
        [T["M"], _f(res['m_min']), _f(res['m_max']), "kNm/m"],
        [T["rot"], _f(res['rot_min'], 4), _f(res['rot_max'], 4), "rad"],
        [T["defl"], _f(float(res['deflection'].min()) * 1000, 1),
         _f(float(res['deflection'].max()) * 1000, 1), "mm"],
    ]))

    def checks_table(src):
        rows = [[T["stress"], f"{_f(src['actual_stress'] / 1000, 1)} MPa",
                 f"{_f(wall.f_allowable / 1000, 1)} MPa", _status(T, src['stress_check_status'])]]
        if src['deflection_check_status'] != "N/A":
            rows.append([T["defl_check"], f"{_f(src['actual_max_deflection'], 1)} mm",
                         f"{_f(src['allowable_deflection'], 1)} mm",
                         _status(T, src['deflection_check_status'])])
        if src.get('vertical_status', "N/A") != "N/A":
            rows.append([T["vertical"], f"{_f(src['vertical_load'], 1)} kN/m",
                         f"{_f(src['vertical_resistance'], 1)} kN/m",
                         _status(T, src['vertical_status'])])
        return _table([T["check"], T["actual"], T["allowable"], T["status"]], rows)

    parts.append(f"<h3>{T['sec_checks']}</h3>")
    parts.append(checks_table(res))

    # ---------------- 3. beam-spring
    if bs is not None:
        parts.append(f"<h2 style='page-break-before:always'>{T['sec_bs']}</h2>")
        parts.append(_kv_table(T, [
            [T["embed_used"], f"{_f(bs['embedment'])} (L = {_f(bs['length'])})", "m"],
            [T["iterations"], str(bs['iterations']), ""],
        ]))
        parts.append(f"<h3>{T['sec_bs_ks']}</h3>")
        parts.append(_table([T["layer"], T["ks_method"], f"{T['ks']} (kN/m³)"],
                            [[_esc(n), L["ks_" + m], _f(k, 0)] for n, k, m in bs['k_s_table']]))
        parts.append(f"<h3>{T['sec_bs_stages']}</h3>")
        rows = []
        for i, s in enumerate(bs['stages'], 1):
            desc = T["excavate" if s['kind'] == 'excavate' else "install"].format(val=s['value'])
            if s['solved']:
                forces = ", ".join(f"T({d:.1f} m) = {t:.0f}" for d, t in sorted(s['anchor_forces'].items()))
                rows.append([i, desc, _f(s['w_max_mm'], 1), forces or "—"])
            else:
                rows.append([i, desc, "—", "—"])
        parts.append(_table([T["stage"], T["description"], f"{T['w_max']} (mm)", T["forces"]], rows,
                            [8, 34, 14, 44]))
        if wall.anchors:
            parts.append(f"<h3>{T['sec_bs_anchor']}</h3>")
            rows = [[i, _f(a['depth']), _f(rep['T_h'], 1), _f(rep['T_axial'], 1), _f(rep['V'], 1),
                     _f(res['anchor_report'][a['depth']]['T_h'], 1)]
                    for i, (a, rep) in enumerate(
                        ((a, bs['anchor_report'][a['depth']]) for a in wall.anchors), 1)]
            parts.append(_table(["#", f"{T['depth']} (m)", f"{T['T_h']} (kN/m)", f"{T['T_ax']} (kN)",
                                 f"{T['V']} (kN/m)", f"LE {T['T_h']} (kN/m)"], rows))
        parts.append(f"<h3>{T['sec_bs_checks']}</h3>")
        parts.append(checks_table(bs))
        parts.append(_table([T["compare"], "LE", "Winkler", T["unit"]], [
            [T["M"], _f(res['m_max_abs'], 1), _f(bs['m_max_abs'], 1), "kNm/m"],
            [T["defl"], _f(res['actual_max_deflection'], 1), _f(bs['actual_max_deflection'], 1), "mm"],
            [T["mob_p"], "—", f"{bs['passive_mobilization']:.0%}", ""],
            [T["mob_a"], "—", f"{bs['active_mobilization']:.0%}", ""],
        ]))

    # ---------------- 4. figures
    parts.append(f"<h2 style='page-break-before:always'>{T['sec_figs']}</h2>")
    for i, key in enumerate([k for k in figures if not k.startswith("study_")], 1):
        title = L.get(f"tab_{key}", key)
        parts.append(f"<p><b>{_esc(title)}</b></p>")
        parts.append(f"<p><img src='{img_src(key)}' width='640'></p>")

    # ---------------- 5/6. warnings & notes
    if res.get('warnings'):
        parts.append(f"<h2>{T['sec_warn']}</h2><ul>")
        parts.extend(f"<li>{_esc(w)}</li>" for w in res['warnings'])
        parts.append("</ul>")
    parts.append(f"<h2>{T['sec_notes']}</h2><ul class='note'>")
    parts.extend(f"<li>{_esc(n)}</li>" for n in T["notes"])
    parts.append("</ul>")
    if study is not None and study.rows:
        parts.extend(_study_section(T, L, study, figures, img_src))
    parts.append("</body></html>")
    return "\n".join(parts)


# ----------------------------------------------------------------------
# exporters
# ----------------------------------------------------------------------
def _all_figures(wall, engine, bs, lang, study):
    figures = render_figures(wall, engine, bs, lang)
    figures.update(render_study_figures(study, lang))
    return figures


def export_html(path: str, wall, engine, bs, lang: str, study=None) -> None:
    figures = _all_figures(wall, engine, bs, lang, study)
    with open(path, 'w', encoding='utf-8') as f:
        f.write(build_html(wall, engine, bs, lang, figures, study=study))


def export_pdf(path: str, wall, engine, bs, lang: str, study=None) -> None:
    """PDF through reportlab; the figures travel as `fig://<key>` references."""
    from . import pdf as pdf_writer

    T = TEXTS[lang]
    info = wall.config['project_info']
    figures = _all_figures(wall, engine, bs, lang, study)
    html_text = build_html(wall, engine, bs, lang, figures,
                           img_src=lambda k: f"fig://{k}", study=study)
    pdf_writer.html_to_pdf(path, html_text, figures,
                           title=info.get('title', T['title']),
                           author=info.get('analyst', ''),
                           footer=f"{info.get('title', '')} — {APP_NAME} v{APP_VERSION}")


def export_docx(path: str, wall, engine, bs, lang: str, study=None) -> None:
    """Word report; requires python-docx."""
    try:
        import docx
        from docx.shared import Inches, Pt
    except ImportError as e:
        raise RuntimeError("python-docx is not installed (pip install python-docx).") from e
    import re

    figures = _all_figures(wall, engine, bs, lang, study)
    html_text = build_html(wall, engine, bs, lang, figures, img_src=lambda k: f"fig://{k}", study=study)
    d = docx.Document()
    d.styles['Normal'].font.size = Pt(10)

    # very small HTML -> docx walker (headings, paragraphs, tables, lists, images)
    tokens = re.split(r"(<h1>.*?</h1>|<h2[^>]*>.*?</h2>|<h3>.*?</h3>|<table[^>]*>.*?</table>|"
                      r"<p[^>]*>.*?</p>|<li>.*?</li>)", html_text, flags=re.S)
    strip = lambda s: html.unescape(re.sub(r"<[^>]+>", "", re.sub(r"<br\s*/?>", "\n", s))).strip()
    for tok in tokens:
        if tok.startswith("<h1>"):
            d.add_heading(strip(tok), level=0)
        elif tok.startswith("<h2"):
            d.add_heading(strip(tok), level=1)
        elif tok.startswith("<h3>"):
            d.add_heading(strip(tok), level=2)
        elif tok.startswith("<li>"):
            d.add_paragraph(strip(tok), style='List Bullet')
        elif tok.startswith("<table"):
            rows = re.findall(r"<tr>(.*?)</tr>", tok, flags=re.S)
            cells = [re.findall(r"<t[hd][^>]*>(.*?)</t[hd]>", r, flags=re.S) for r in rows]
            if cells:
                table = d.add_table(rows=len(cells), cols=len(cells[0]))
                table.style = 'Light Grid Accent 1'
                for i, row in enumerate(cells):
                    for j, c in enumerate(row):
                        if j < len(table.columns):
                            table.cell(i, j).text = strip(c)
        elif tok.startswith("<p"):
            m = re.search(r"src='fig://([a-z_]+)'", tok)
            if m:
                d.add_picture(io.BytesIO(figures[m.group(1)]), width=Inches(6.3))
            else:
                text = strip(tok)
                if text:
                    d.add_paragraph(text)
    d.save(path)


def export_report(path: str, wall, engine, bs, lang: str, study=None) -> None:
    ext = path.lower().rsplit('.', 1)[-1]
    if ext == 'pdf':
        export_pdf(path, wall, engine, bs, lang, study)
    elif ext == 'docx':
        export_docx(path, wall, engine, bs, lang, study)
    else:
        export_html(path if ext in ('html', 'htm') else path + '.html', wall, engine, bs, lang, study)
