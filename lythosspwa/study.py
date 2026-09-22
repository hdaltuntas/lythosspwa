"""
Parametric and probabilistic (reliability) studies for Lythos SPWA.

A study is defined by a list of variables, each addressing one input of
the project configuration by a dotted path (e.g. ``soil_profile.0.phi``,
``geometry.excavation_depth_H``, ``anchors.1.prestress``), and either

* a **range** (min, max)             -> parametric / sensitivity study, or
* a **distribution** (normal, lognormal, uniform; mean, CoV) -> reliability.

Sampling methods
    oat   one-at-a-time sweep over each range variable (others at base value)
    grid  full factorial grid over range variables
    lhs   Latin hypercube (ranges -> uniform, distributions -> inverse CDF)
    mc    plain Monte Carlo
Optional correlation between distribution variables (Gaussian copula).

Each sample is analysed with the LE engine (and optionally the beam-spring
model); the outputs are collected in a flat table. Post-processing gives
summary statistics, Spearman rank sensitivities, standardised regression
coefficients and, for the limit states, failure probabilities with
confidence intervals and reliability indices.

The runner works in parallel (process pool) and supports a progress
callback and cooperative cancellation.
"""

import copy
import csv
import math
import os
from concurrent.futures import ProcessPoolExecutor, as_completed
from typing import Any, Callable, Dict, List, Optional, Sequence, Tuple

import numpy as np
from scipy import stats
from scipy.stats import qmc

from .analysis_engine import AnalysisEngine, RetainingWall
from .beam_spring import BeamSpringAnalysis

# ----------------------------------------------------------------------
# variables
# ----------------------------------------------------------------------
DISTRIBUTIONS = ["normal", "lognormal", "uniform"]
METHODS = ["oat", "grid", "lhs", "mc"]

# outputs collected for every sample (key, label-key, unit)
OUTPUTS = [
    ("d_req", "d_req", "m"), ("d_design", "d_design", "m"),
    ("M_le", "M_le", "kNm/m"), ("sigma_le", "sigma_le", "MPa"),
    ("defl_le", "defl_le", "mm"), ("T_le", "T_le", "kN/m"),
    ("M_bs", "M_bs", "kNm/m"), ("sigma_bs", "sigma_bs", "MPa"),
    ("defl_bs", "defl_bs", "mm"), ("T_bs", "T_bs", "kN/m"),
    ("mob_p", "mob_p", "-"),
]

# limit states: g = capacity - demand (failure when g < 0)
LIMIT_STATES = {
    "stress_le": ("sigma_le", "f_allow"),
    "defl_le": ("defl_le", "defl_allow"),
    "stress_bs": ("sigma_bs", "f_allow"),
    "defl_bs": ("defl_bs", "defl_allow"),
}


def available_variables(cfg: Dict[str, Any]) -> List[Tuple[str, str]]:
    """(path, group) pairs of the inputs that can be varied for this project."""
    out = [
        ("geometry.excavation_depth_H", "geometry"),
        ("geometry.wall_friction_delta", "geometry"),
        ("geometry.backfill_slope_beta", "geometry"),
        ("loads.surcharge_load", "loads"),
        ("loads.water_level_active", "loads"),
        ("loads.water_level_passive", "loads"),
        ("analysis_options.kh", "seismic"),
        ("factors.embedment_increase_factor", "factors"),
    ]
    for i, lay in enumerate(cfg.get("soil_profile", [])):
        for key in ("phi", "cohesion", "gamma", "gamma_sat", "thickness", "k_s", "E_M"):
            out.append((f"soil_profile.{i}.{key}", lay.get("name", f"layer {i + 1}")))
    for i, _ in enumerate(cfg.get("analysis_options", {}).get("anchors", [])):
        for key in ("depth", "prestress", "EA", "free_length", "spacing", "angle"):
            out.append((f"analysis_options.anchors.{i}.{key}", f"anchor {i + 1}"))
    return out


def pretty_label(cfg: Dict[str, Any], path: str, lang: Dict[str, str]) -> str:
    """Human-readable variable name, e.g. 'Sandy Gravel · φ' or 'Anchor 2 · P₀'."""
    parts = path.split(".")
    key = parts[-1]
    name = lang.get(f"var_{key}", key)
    if parts[0] == "soil_profile":
        layer = cfg["soil_profile"][int(parts[1])].get("name", f"layer {int(parts[1]) + 1}")
        return f"{layer} · {name}"
    if parts[0] == "analysis_options" and parts[1] == "anchors":
        return f"{lang.get('grp_anchor', 'Anchor {i}').format(i=int(parts[2]) + 1)} · {name}"
    group = {"geometry": "grp_geometry", "loads": "grp_loads", "factors": "grp_factors",
             "analysis_options": "grp_seismic"}.get(parts[0])
    return f"{lang.get(group, parts[0])} · {name}" if group else name


def get_value(cfg: Dict[str, Any], path: str) -> Any:
    node = cfg
    for part in path.split("."):
        node = node[int(part)] if isinstance(node, list) else node[part]
    return node


def set_value(cfg: Dict[str, Any], path: str, value: Any) -> None:
    parts = path.split(".")
    node = cfg
    for part in parts[:-1]:
        node = node[int(part)] if isinstance(node, list) else node[part]
    last = parts[-1]
    if isinstance(node, list):
        node[int(last)] = value
    else:
        node[last] = value


class StudyVariable:
    """One varied input. mode = 'range' (min/max) or 'dist' (distribution)."""

    def __init__(self, path: str, mode: str = "range", vmin: float = 0.0, vmax: float = 1.0,
                 dist: str = "normal", mean: float = 0.0, cov: float = 0.1,
                 n_points: int = 5, label: Optional[str] = None):
        self.path, self.mode = path, mode
        self.vmin, self.vmax = float(vmin), float(vmax)
        self.dist, self.mean, self.cov = dist, float(mean), float(cov)
        self.n_points = int(n_points)
        self.label = label or path
        self.base: Optional[float] = None    # project value, set by Study
        if mode not in ("range", "dist"):
            raise ValueError(f"{path}: mode must be 'range' or 'dist'")
        if mode == "range" and self.vmax <= self.vmin:
            raise ValueError(f"{path}: max must be > min")
        if mode == "dist":
            if dist not in DISTRIBUTIONS:
                raise ValueError(f"{path}: unknown distribution '{dist}'")
            if self.cov <= 0 or (dist != "uniform" and self.mean <= 0 and dist == "lognormal"):
                raise ValueError(f"{path}: CoV must be > 0 (and mean > 0 for lognormal)")

    # -- inverse CDF / value from a uniform quantile u in (0, 1)
    def ppf(self, u: np.ndarray) -> np.ndarray:
        u = np.clip(u, 1e-9, 1 - 1e-9)
        if self.mode == "range":
            return self.vmin + (self.vmax - self.vmin) * u
        std = self.mean * self.cov
        if self.dist == "normal":
            return stats.norm.ppf(u, loc=self.mean, scale=std)
        if self.dist == "lognormal":
            s2 = math.log(1.0 + self.cov ** 2)
            mu = math.log(self.mean) - 0.5 * s2
            return np.exp(stats.norm.ppf(u, loc=mu, scale=math.sqrt(s2)))
        # uniform with given mean and CoV: half-width = sqrt(3)*std
        hw = math.sqrt(3.0) * std
        return self.mean - hw + 2.0 * hw * u

    def to_dict(self) -> Dict[str, Any]:
        return dict(path=self.path, mode=self.mode, min=self.vmin, max=self.vmax, dist=self.dist,
                    mean=self.mean, cov=self.cov, n_points=self.n_points, label=self.label)

    @classmethod
    def from_dict(cls, d: Dict[str, Any]) -> "StudyVariable":
        return cls(d["path"], d.get("mode", "range"), d.get("min", 0.0), d.get("max", 1.0),
                   d.get("dist", "normal"), d.get("mean", 0.0), d.get("cov", 0.1),
                   d.get("n_points", 5), d.get("label"))


# ----------------------------------------------------------------------
# sampling
# ----------------------------------------------------------------------
def sample(variables: Sequence[StudyVariable], method: str, n: int = 100,
           seed: Optional[int] = 0,
           correlation: Optional[Dict[Tuple[str, str], float]] = None) -> Tuple[np.ndarray, List[int]]:
    """
    Returns (X, sweep_index): X is an (N, k) matrix of input values, one
    column per variable. For 'oat', sweep_index[i] is the index of the
    variable swept in sample i (others at their base = mid-range / mean).
    """
    k = len(variables)
    if k == 0:
        raise ValueError("No study variables defined.")
    rng = np.random.default_rng(seed)
    base = np.array([v.base if v.base is not None else
                     (v.ppf(np.array([0.5]))[0] if v.mode == "dist" else 0.5 * (v.vmin + v.vmax))
                     for v in variables])

    if method == "oat":
        rows, idx = [], []
        for j, v in enumerate(variables):
            pts = v.ppf(np.linspace(0.0, 1.0, v.n_points)) if v.mode == "range" \
                else v.ppf(np.linspace(0.02, 0.98, v.n_points))
            for p in pts:
                r = base.copy(); r[j] = p
                rows.append(r); idx.append(j)
        return np.array(rows), idx

    if method == "grid":
        axes = [v.ppf(np.linspace(0.0, 1.0, v.n_points)) if v.mode == "range"
                else v.ppf(np.linspace(0.02, 0.98, v.n_points)) for v in variables]
        mesh = np.meshgrid(*axes, indexing="ij")
        X = np.column_stack([m.ravel() for m in mesh])
        return X, [-1] * len(X)

    if method == "lhs":
        u = qmc.LatinHypercube(d=k, seed=seed).random(n)
    elif method == "mc":
        u = rng.random((n, k))
    else:
        raise ValueError(f"Unknown sampling method '{method}'")

    if correlation:
        # Gaussian copula: correlate the standard-normal scores, map back to uniforms
        R = np.eye(k)
        names = [v.path for v in variables]
        for (a, b), rho in correlation.items():
            if a in names and b in names:
                i, j = names.index(a), names.index(b)
                R[i, j] = R[j, i] = rho
        Lc = np.linalg.cholesky(R)
        z = stats.norm.ppf(np.clip(u, 1e-9, 1 - 1e-9)) @ Lc.T
        u = stats.norm.cdf(z)

    X = np.column_stack([v.ppf(u[:, j]) for j, v in enumerate(variables)])
    return X, [-1] * len(X)


def _worker_init() -> None:
    """Keep every worker process single-threaded (BLAS/OpenMP), otherwise the
    processes fight for the cores and the pool becomes slower than serial."""
    for var in ("OMP_NUM_THREADS", "OPENBLAS_NUM_THREADS", "MKL_NUM_THREADS",
                "NUMEXPR_NUM_THREADS"):
        os.environ[var] = "1"
    try:
        from threadpoolctl import threadpool_limits
        threadpool_limits(1)
    except Exception:
        pass


# ----------------------------------------------------------------------
# single evaluation (top-level so it can be pickled for the process pool)
# ----------------------------------------------------------------------
def evaluate(base_cfg: Dict[str, Any], paths: Sequence[str], values: Sequence[float],
             run_bs: bool, unfactored: bool) -> Dict[str, Any]:
    cfg = copy.deepcopy(base_cfg)
    for p, v in zip(paths, values):
        set_value(cfg, p, float(v))
    if unfactored:
        cfg['factors'].update({'FS_friction_angle': 1.0, 'FS_cohesion': 1.0, 'FS_bending': 1.0})
    row: Dict[str, Any] = {p: float(v) for p, v in zip(paths, values)}
    row.update({k: float('nan') for k, _, _ in OUTPUTS})
    row.update({'ok': False, 'ok_bs': False, 'error': ''})
    try:
        wall = RetainingWall(cfg)
        eng = AnalysisEngine(wall)
        eng.run()
        r = eng.results
        row.update({
            'd_req': eng.d_required, 'd_design': eng.d_design,
            'M_le': r['m_max_abs'], 'sigma_le': r['actual_stress'] / 1000.0,
            'defl_le': r['actual_max_deflection'],
            'T_le': max(eng.t_anchors.values()) if eng.t_anchors else 0.0,
            'f_allow': wall.f_allowable / 1000.0,
            'defl_allow': r['allowable_deflection'],
            'ok': True,
        })
        if run_bs:
            try:
                bs = BeamSpringAnalysis(wall, eng).run()
                row.update({
                    'M_bs': bs['m_max_abs'], 'sigma_bs': bs['actual_stress'] / 1000.0,
                    'defl_bs': bs['actual_max_deflection'],
                    'T_bs': max(bs['anchor_forces'].values()) if bs['anchor_forces'] else 0.0,
                    'mob_p': bs['passive_mobilization'], 'ok_bs': True,
                })
            except Exception as e:          # unstable wall etc.
                row['error'] = f"BS: {e}"
    except Exception as e:
        row['error'] = f"{type(e).__name__}: {e}"
    return row


# ----------------------------------------------------------------------
# runner
# ----------------------------------------------------------------------
class Study:
    def __init__(self, base_cfg: Dict[str, Any], variables: Sequence[StudyVariable],
                 method: str = "lhs", n: int = 100, run_bs: bool = True,
                 unfactored: bool = False, seed: Optional[int] = 0,
                 correlation: Optional[Dict[Tuple[str, str], float]] = None,
                 workers: Optional[int] = None):
        self.base_cfg = base_cfg
        self.variables = list(variables)
        for v in self.variables:
            if v.base is None:
                try:
                    val = float(get_value(base_cfg, v.path))
                except (KeyError, IndexError, TypeError, ValueError):
                    val = None
                if val is not None and (v.mode == "dist" or v.vmin <= val <= v.vmax):
                    v.base = val
        self.method, self.n = method, int(n)
        self.run_bs, self.unfactored, self.seed = run_bs, unfactored, seed
        self.correlation = correlation or {}
        self.workers = workers if workers is not None else max(1, (os.cpu_count() or 2) - 1)
        self.X: Optional[np.ndarray] = None
        self.sweep_index: List[int] = []
        self.rows: List[Dict[str, Any]] = []
        self.cancelled = False
        self.summary: Dict[str, Any] = {}

    @property
    def paths(self) -> List[str]:
        return [v.path for v in self.variables]

    def run(self, progress: Optional[Callable[[int, int], None]] = None,
            is_cancelled: Optional[Callable[[], bool]] = None) -> List[Dict[str, Any]]:
        self.X, self.sweep_index = sample(self.variables, self.method, self.n,
                                          self.seed, self.correlation)
        n_total = len(self.X)
        self.rows = [None] * n_total
        args = (self.base_cfg, self.paths)

        def store(i, row):
            row['sample'] = i
            row['sweep_var'] = self.variables[self.sweep_index[i]].path if self.sweep_index[i] >= 0 else ''
            self.rows[i] = row

        done = 0
        if self.workers > 1 and n_total >= 8:
            with ProcessPoolExecutor(max_workers=self.workers, initializer=_worker_init) as pool:
                futures = {pool.submit(evaluate, *args, self.X[i], self.run_bs, self.unfactored): i
                           for i in range(n_total)}
                for fut in as_completed(futures):
                    i = futures[fut]
                    store(i, fut.result())
                    done += 1
                    if progress:
                        progress(done, n_total)
                    if is_cancelled and is_cancelled():
                        self.cancelled = True
                        for f in futures:
                            f.cancel()
                        break
        else:
            for i in range(n_total):
                store(i, evaluate(*args, self.X[i], self.run_bs, self.unfactored))
                done += 1
                if progress:
                    progress(done, n_total)
                if is_cancelled and is_cancelled():
                    self.cancelled = True
                    break
        self.rows = [r for r in self.rows if r is not None]
        self.summary = summarize(self)
        return self.rows

    # -- serialisation
    def spec_dict(self) -> Dict[str, Any]:
        return {'method': self.method, 'n': self.n, 'run_bs': self.run_bs,
                'unfactored': self.unfactored, 'seed': self.seed,
                'variables': [v.to_dict() for v in self.variables],
                'correlation': [[a, b, r] for (a, b), r in self.correlation.items()]}


# ----------------------------------------------------------------------
# post-processing
# ----------------------------------------------------------------------
def _column(rows: List[Dict[str, Any]], key: str) -> np.ndarray:
    return np.array([r.get(key, float('nan')) for r in rows], dtype=float)


def reliability(g: np.ndarray) -> Dict[str, float]:
    """Failure probability from samples of the margin g (fail if g < 0)."""
    g = g[np.isfinite(g)]
    n = len(g)
    if n == 0:
        return {'n': 0, 'n_fail': 0, 'pf': float('nan'), 'pf_lo': float('nan'),
                'pf_hi': float('nan'), 'beta': float('nan'), 'beta_fosm': float('nan')}
    nf = int(np.sum(g < 0))
    pf = nf / n
    # Wilson 95 % interval
    z = 1.96
    denom = 1 + z * z / n
    centre = (pf + z * z / (2 * n)) / denom
    half = z * math.sqrt(pf * (1 - pf) / n + z * z / (4 * n * n)) / denom
    lo, hi = max(0.0, centre - half), min(1.0, centre + half)
    beta = float(-stats.norm.ppf(pf)) if 0 < pf < 1 else (float('inf') if pf == 0 else float('-inf'))
    sg = float(np.std(g, ddof=1)) if n > 1 else 0.0
    beta_fosm = float(np.mean(g) / sg) if sg > 0 else float('inf')
    return {'n': n, 'n_fail': nf, 'pf': pf, 'pf_lo': lo, 'pf_hi': hi,
            'beta': beta, 'beta_fosm': beta_fosm}


def summarize(study: "Study") -> Dict[str, Any]:
    rows = study.rows
    ok = [r for r in rows if r['ok']]
    out: Dict[str, Any] = {'n_total': len(rows), 'n_ok': len(ok),
                           'n_ok_bs': sum(1 for r in ok if r['ok_bs']),
                           'stats': {}, 'spearman': {}, 'src': {}, 'reliability': {},
                           'oat': {}}
    if not ok:
        return out
    X = np.column_stack([_column(ok, p) for p in study.paths])
    for key, _, _ in OUTPUTS:
        y = _column(ok, key)
        m = np.isfinite(y)
        if m.sum() < 2:
            continue
        yy = y[m]
        out['stats'][key] = {'mean': float(np.mean(yy)), 'std': float(np.std(yy, ddof=1)),
                             'min': float(np.min(yy)), 'max': float(np.max(yy)),
                             'p5': float(np.percentile(yy, 5)), 'p50': float(np.percentile(yy, 50)),
                             'p95': float(np.percentile(yy, 95)), 'n': int(m.sum())}
        if study.method in ("lhs", "mc", "grid") and len(study.paths) >= 1 and m.sum() >= 5:
            sp, src = {}, {}
            Xm = X[m]
            for j, p in enumerate(study.paths):
                if np.std(Xm[:, j]) > 0 and np.std(yy) > 0:
                    sp[p] = float(stats.spearmanr(Xm[:, j], yy).statistic)
            # standardised regression coefficients (linear surrogate)
            if len(study.paths) < m.sum() - 1:
                A = np.column_stack([np.ones(len(yy)), Xm])
                try:
                    coef, *_ = np.linalg.lstsq(A, yy, rcond=None)
                    for j, p in enumerate(study.paths):
                        sx = np.std(Xm[:, j], ddof=1)
                        src[p] = float(coef[j + 1] * sx / np.std(yy, ddof=1)) if sx > 0 else 0.0
                except np.linalg.LinAlgError:
                    pass
            out['spearman'][key] = sp
            out['src'][key] = src
    # limit states (only meaningful for random sampling)
    for name, (demand, capacity) in (LIMIT_STATES.items() if study.method in ("lhs", "mc") else []):
        d, c = _column(ok, demand), _column(ok, capacity)
        g = c - d
        if np.isfinite(g).sum() >= 2 and np.isfinite(c).all():
            out['reliability'][name] = reliability(g)
    # one-at-a-time sweeps: per variable, sorted (x, outputs)
    if study.method == "oat":
        for j, v in enumerate(study.variables):
            sub = [r for r in ok if r['sweep_var'] == v.path]
            if not sub:
                continue
            xs = _column(sub, v.path)
            order = np.argsort(xs)
            out['oat'][v.path] = {'x': xs[order].tolist(),
                                  **{k: _column(sub, k)[order].tolist() for k, _, _ in OUTPUTS}}
    return out


# ----------------------------------------------------------------------
# export
# ----------------------------------------------------------------------
def table_columns(study: "Study") -> List[str]:
    return ['sample', 'sweep_var'] + study.paths + [k for k, _, _ in OUTPUTS] + \
        ['f_allow', 'defl_allow', 'ok', 'ok_bs', 'error']


def to_csv(study: "Study", path: str) -> None:
    cols = table_columns(study)
    with open(path, 'w', newline='', encoding='utf-8') as f:
        w = csv.DictWriter(f, fieldnames=cols, extrasaction='ignore')
        w.writeheader()
        for r in study.rows:
            w.writerow({c: r.get(c, '') for c in cols})


def to_xlsx(study: "Study", path: str) -> None:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font
    except ImportError as e:
        raise RuntimeError("openpyxl is not installed (pip install openpyxl).") from e
    wb = Workbook()
    ws = wb.active
    ws.title = "samples"
    cols = table_columns(study)
    ws.append(cols)
    for c in ws[1]:
        c.font = Font(bold=True)
    for r in study.rows:
        ws.append([r.get(c, '') if not (isinstance(r.get(c), float) and math.isnan(r.get(c)))
                   else None for c in cols])
    ws2 = wb.create_sheet("summary")
    s = study.summary
    ws2.append(["output", "n", "mean", "std", "min", "p5", "p50", "p95", "max"])
    for k, st in s.get('stats', {}).items():
        ws2.append([k, st['n'], st['mean'], st['std'], st['min'], st['p5'], st['p50'], st['p95'], st['max']])
    ws2.append([])
    ws2.append(["limit state", "n", "n_fail", "Pf", "Pf 95% lo", "Pf 95% hi", "beta", "beta_FOSM"])
    for k, rl in s.get('reliability', {}).items():
        ws2.append([k, rl['n'], rl['n_fail'], rl['pf'], rl['pf_lo'], rl['pf_hi'], rl['beta'], rl['beta_fosm']])
    ws2.append([])
    ws2.append(["Spearman rho", *study.paths])
    for k, sp in s.get('spearman', {}).items():
        ws2.append([k, *[sp.get(p, None) for p in study.paths]])
    ws3 = wb.create_sheet("variables")
    ws3.append(["path", "mode", "min", "max", "dist", "mean", "cov", "n_points"])
    for v in study.variables:
        d = v.to_dict()
        ws3.append([d['path'], d['mode'], d['min'], d['max'], d['dist'], d['mean'], d['cov'], d['n_points']])
    wb.save(path)
